"""Real file conversions; all fixtures are made here without downloads."""
import io
import os
import zipfile
from markitdown import MarkItDown
from openpyxl import Workbook
from pptx import Presentation

os.setgroups([])
os.setgid(1001)
os.setuid(1001)
text = 'OpenCrane document proof'
stream = f'BT /F1 18 Tf 20 80 Td ({text}) Tj ET'.encode()
objects = [b'<< /Type /Catalog /Pages 2 0 R >>',
           b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
           b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 120] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
           b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
           b'<< /Length %d >>\nstream\n' % len(stream) + stream + b'\nendstream']
pdf = b'%PDF-1.4\n'
offsets = [0]
for i, obj in enumerate(objects, 1):
    offsets.append(len(pdf))
    pdf += f'{i} 0 obj\n'.encode() + obj + b'\nendobj\n'
xref = len(pdf)
pdf += b'xref\n0 6\n0000000000 65535 f \n'
pdf += b''.join(f'{n:010} 00000 n \n'.encode() for n in offsets[1:])
pdf += f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode()
converter = MarkItDown()
assert text in converter.convert_stream(io.BytesIO(pdf), file_extension='.pdf').text_content
with zipfile.ZipFile('/workspace/sample.docx', 'w') as doc:
    doc.writestr('[Content_Types].xml', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    doc.writestr('word/document.xml', f'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:body></w:document>')
book = Workbook()
book.active['A1'] = text
book.save('/workspace/sample.xlsx')
presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[0])
slide.shapes.title.text = text
presentation.save('/workspace/sample.pptx')
for extension in ['docx', 'xlsx', 'pptx']:
    assert text in converter.convert('/workspace/sample.' + extension).text_content
print('DOCUMENT_PASS')
